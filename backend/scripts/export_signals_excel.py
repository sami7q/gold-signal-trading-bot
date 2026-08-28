"""Export every entry signal the Telegram channel posted, from --since to now,
to a 3-sheet Excel workbook (explicit user request: precise, complete, no
silent gaps, Arabic AND English entry signals both caught).

Reuses, rather than reimplementing:
  - `build_channel_archive._fetch_channel_messages` for Telegram history
    fetching -- same Telethon session, same full-pagination reverse-order
    iteration the archive builder already uses and has been validated
    against real channel history with.
  - `app.signals.parser` (normalize_text, detect_kind, is_gold_message,
    extract_stop_loss, the STOP_LOSS regex) -- the exact same signal
    detection the LIVE bot uses, so "confirmed signal" here means "the bot
    would have classified this as an entry signal", not a second,
    independently-drifting interpretation.

Arabic support status (checked against the current parser, not assumed):
  app/signals/parser.py ALREADY recognizes "شراء الذهب"/"بيع الذهب" as
  BUY/SELL market entries (added in an earlier phase of this project -- this
  is not a fresh fix). What it does NOT have is any Arabic pattern for
  pending/conditional orders (BUY LIMIT/STOP equivalents) -- confirmed via
  its own docstring: no Arabic LIMIT/STOP phrasing had been observed in
  channel history at the time, so none was guessed. That gap is exactly what
  the "Review Needed" bucket below exists to catch, rather than silently
  missing it.

Bucketing -- every scanned message lands in EXACTLY one of three buckets
(enforced by an assertion at the end, not just hoped for):
  1. CONFIRMED SIGNAL    -- parser classifies it ENTRY_SIGNAL (a kind was
     detected AND gold was mentioned). Six possible order types.
  2. REVIEW NEEDED       -- not classified as ENTRY_SIGNAL, but mentions
     gold plus buy/sell/pending vocabulary (Arabic or English) that didn't
     match the parser's confirmed patterns -- most importantly Arabic
     pending-order phrasing per the gap above. Never silently dropped.
  3. CONFIRMED NON-SIGNAL -- everything else (SL/TP updates, results,
     unrelated commentary, messages with no gold/order-type vocabulary at all).

Usage (from backend/):
    .venv/Scripts/python.exe scripts/export_signals_excel.py \
        [--since 2026-01-01] [--channel-id ID] [--out signals_export.xlsx]
"""

from __future__ import annotations

import argparse
import asyncio
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import UTC, date, datetime, timedelta
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import build_channel_archive as archive  # noqa: E402 -- reuse its Telegram fetch, not a copy
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Font  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from app.core.config import get_settings  # noqa: E402
from app.core.constants import SignalKind  # noqa: E402
from app.signals.parser import (  # noqa: E402
    _STOP_LOSS_PATTERN,  # reused to strip SL-update mentions before the review heuristic
    detect_kind,
    is_gold_message,
    normalize_text,
)

_KIND_TO_ORDER_TYPE: dict[SignalKind, str] = {
    SignalKind.BUY: "BUY_MARKET",
    SignalKind.SELL: "SELL_MARKET",
    SignalKind.BUY_LIMIT: "BUY_LIMIT",
    SignalKind.SELL_LIMIT: "SELL_LIMIT",
    SignalKind.BUY_STOP: "BUY_STOP",
    SignalKind.SELL_STOP: "SELL_STOP",
}

_ORDER_TYPE_COLUMNS = [
    "BUY_MARKET",
    "SELL_MARKET",
    "BUY_LIMIT",
    "SELL_LIMIT",
    "BUY_STOP",
    "SELL_STOP",
]

_AR_PENDING_WORDS = ("معلق", "مؤجل")
_AR_DIRECTION_WORDS = ("شراء", "بيع")
_EN_TRIGGER_PATTERN = re.compile(r"\b(BUY|SELL|LIMIT|STOP|PENDING)\b")

#: Deliberately broader than `is_gold_message` (which requires the parser's
#: exact `\bالذهب\b`-anchored token). Arabic attaches grammatical prefixes
#: directly to the noun with no space -- "للذهب" ("for/to the gold"),
#: "بالذهب" ("with the gold"), "كالذهب" ("like gold") -- and none of those
#: contain "الذهب" as a *word-bounded* substring (there's no boundary
#: between the attached prefix letter and the alef of "ال"), so the strict
#: parser check silently misses them. A bare root-substring check on "ذهب"
#: (plus the English aliases) catches every such form. This is a "should I
#: even look closer at this message" gate, not a final classification --
#: over-matching here only means an extra Review Needed row, never a
#: silently dropped signal.
_GOLD_ROOT_PATTERN = re.compile(r"ذهب|GOLD|XAU")


def _mentions_gold_loosely(normalized_text: str) -> bool:
    return bool(_GOLD_ROOT_PATTERN.search(normalized_text))


def _looks_signal_ish(normalized_text: str) -> tuple[bool, str]:
    """Second pass over messages the strict parser did NOT classify as
    ENTRY_SIGNAL: could this still plausibly be a missed signal? Returns
    (flag, reason). Deliberately noisier than the parser's own patterns --
    a false positive here just costs a human one extra glance in the
    Review Needed sheet; a false negative means a real signal silently
    vanishes, which is the one outcome this task forbids."""
    if not _mentions_gold_loosely(normalized_text):
        return False, ""

    if any(w in normalized_text for w in _AR_PENDING_WORDS):
        return True, (
            "Gold + Arabic pending-order vocabulary (معلق/مؤجل) — "
            "parser has no Arabic LIMIT/STOP pattern"
        )

    if any(w in normalized_text for w in _AR_DIRECTION_WORDS):
        return True, "Gold + شراء/بيع present but not in the confirmed adjacent 'X الذهب' pattern"

    # Strip legitimate "SL 4623" / "STOP LOSS 4623" mentions first so a
    # plain Stop-Loss update doesn't false-positive on the word STOP.
    text_without_sl = _STOP_LOSS_PATTERN.sub(" ", normalized_text)
    match = _EN_TRIGGER_PATTERN.search(text_without_sl)
    if match:
        return True, (
            f"Gold + '{match.group(1)}' present but not classified as a confirmed entry signal"
        )

    return False, ""


@dataclass(slots=True)
class SignalRow:
    message_id: str
    day: date
    entry_time: str  # HH:MM UTC
    order_type: str
    raw_text: str
    day_sequence_number: int = 0


@dataclass(slots=True)
class ReviewRow:
    message_id: str
    day: date
    time: str
    raw_text: str
    reason: str


@dataclass(slots=True)
class ExportResult:
    since: datetime
    until: datetime
    total_scanned: int = 0
    signals: list[SignalRow] = field(default_factory=list)
    review: list[ReviewRow] = field(default_factory=list)
    non_signal_count: int = 0


def classify_all(messages: list, since: datetime, until: datetime) -> ExportResult:
    result = ExportResult(since=since, until=until, total_scanned=len(messages))
    day_counters: dict[date, int] = defaultdict(int)

    for msg in messages:
        normalized = normalize_text(msg.text)
        kind = detect_kind(normalized)
        gold = is_gold_message(normalized)
        day = msg.sent_at.astimezone(UTC).date()
        time_str = msg.sent_at.astimezone(UTC).strftime("%H:%M")

        if kind is not None and gold and kind in _KIND_TO_ORDER_TYPE:
            day_counters[day] += 1
            result.signals.append(
                SignalRow(
                    message_id=msg.message_id,
                    day=day,
                    entry_time=time_str,
                    order_type=_KIND_TO_ORDER_TYPE[kind],
                    raw_text=msg.text,
                    day_sequence_number=day_counters[day],
                )
            )
            continue

        flagged, reason = _looks_signal_ish(normalized)
        if flagged:
            result.review.append(
                ReviewRow(
                    message_id=msg.message_id,
                    day=day,
                    time=time_str,
                    raw_text=msg.text,
                    reason=reason,
                )
            )
            continue

        result.non_signal_count += 1

    # Completeness proof -- structural, not just reported: every message hit
    # exactly one branch above, so this can only fail if the bucketing logic
    # itself has a gap (a message counted in more than one place, or in none).
    accounted = len(result.signals) + len(result.review) + result.non_signal_count
    assert accounted == result.total_scanned, (
        f"Bucket counts ({accounted}) do not equal total messages scanned "
        f"({result.total_scanned}) -- this run is INVALID, do not trust the output."
    )
    return result


def _daily_summary_rows(result: ExportResult) -> list[tuple]:
    by_day: dict[date, dict[str, int]] = defaultdict(lambda: dict.fromkeys(_ORDER_TYPE_COLUMNS, 0))
    for row in result.signals:
        by_day[row.day][row.order_type] += 1

    start_day = result.since.astimezone(UTC).date()
    end_day = result.until.astimezone(UTC).date()
    rows = []
    day = start_day
    while day <= end_day:
        counts = by_day.get(day, dict.fromkeys(_ORDER_TYPE_COLUMNS, 0))
        total = sum(counts.values())
        rows.append((day, total, *[counts[c] for c in _ORDER_TYPE_COLUMNS]))
        day += timedelta(days=1)
    return rows


def write_excel(result: ExportResult, out_path: Path) -> None:
    wb = Workbook()
    bold = Font(bold=True)

    ws = wb.active
    ws.title = "Signals"
    headers = ["Date", "Day Sequence #", "Entry Time (UTC)", "Order Type", "Message ID", "Raw Text"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = bold
    for sig_row in sorted(
        result.signals, key=lambda r: (r.day, r.entry_time, r.day_sequence_number)
    ):
        ws.append(
            [
                sig_row.day.isoformat(),
                sig_row.day_sequence_number,
                sig_row.entry_time,
                sig_row.order_type,
                sig_row.message_id,
                sig_row.raw_text,
            ]
        )
    for col, width in zip("ABCDEF", (12, 14, 15, 12, 14, 80), strict=False):
        ws.column_dimensions[col].width = width

    ws2 = wb.create_sheet("Daily Summary")
    headers2 = [
        "Date",
        "Total Signals",
        "Buy Market Count",
        "Sell Market Count",
        "Buy Limit Count",
        "Sell Limit Count",
        "Buy Stop Count",
        "Sell Stop Count",
    ]
    ws2.append(headers2)
    for cell in ws2[1]:
        cell.font = bold
    for summary_row in _daily_summary_rows(result):
        ws2.append([summary_row[0].isoformat(), *summary_row[1:]])
    for i, width in enumerate((12, 14, 16, 16, 15, 15, 15, 15), start=1):
        ws2.column_dimensions[get_column_letter(i)].width = width

    ws3 = wb.create_sheet("Review Needed")
    headers3 = ["Message ID", "Date", "Time (UTC)", "Raw Text", "Reason Flagged"]
    ws3.append(headers3)
    for cell in ws3[1]:
        cell.font = bold
    for rev_row in sorted(result.review, key=lambda r: (r.day, r.time)):
        ws3.append(
            [
                rev_row.message_id,
                rev_row.day.isoformat(),
                rev_row.time,
                rev_row.raw_text,
                rev_row.reason,
            ]
        )
    for col, width in zip("ABCDE", (14, 12, 12, 80, 60), strict=False):
        ws3.column_dimensions[col].width = width

    wb.save(out_path)


async def _run(channel_id: int, since: datetime, out_path: Path) -> None:
    messages = await archive._fetch_channel_messages(channel_id, since)
    until = datetime.now(UTC)
    result = classify_all(messages, since, until)
    write_excel(result, out_path)

    zero_signal_days = [
        row[0] for row in _daily_summary_rows(result) if row[1] == 0
    ]

    print("\n" + "=" * 70)
    print("FINAL REPORT")
    print("=" * 70)
    print(f"Date range covered:        {since.date()} .. {until.date()} (UTC)")
    print(f"Total messages scanned:    {result.total_scanned}")
    print(f"Confirmed signals:         {len(result.signals)}")
    print(f"Confirmed non-signals:     {result.non_signal_count}")
    print(f"Flagged for review:        {len(result.review)}")
    accounted = len(result.signals) + result.non_signal_count + len(result.review)
    print(f"Sum check: {len(result.signals)} + {result.non_signal_count} + {len(result.review)} "
          f"= {accounted}  ->  {'MATCHES' if accounted == result.total_scanned else 'MISMATCH'} "
          f"total scanned ({result.total_scanned})")
    print(f"\nDates in range with ZERO confirmed signals ({len(zero_signal_days)}):")
    if zero_signal_days:
        for d in zero_signal_days:
            print(f"  - {d.isoformat()}")
    else:
        print("  (none -- every day in range has at least one confirmed signal)")
    print(f"\nWorkbook written to: {out_path.resolve()}")
    print("=" * 70)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--since", default="2026-01-01", help="YYYY-MM-DD, default 2026-01-01")
    parser.add_argument(
        "--channel-id", type=int, default=None, help="Defaults to TELEGRAM_ALLOWED_CHANNEL_ID"
    )
    parser.add_argument("--out", default="signals_export.xlsx", help="Output .xlsx path")
    args = parser.parse_args()

    settings = get_settings()
    channel_id = args.channel_id or int(settings.telegram_allowed_channel_id)
    since = datetime.strptime(args.since, "%Y-%m-%d").replace(tzinfo=UTC)

    print(f"Fetching channel {channel_id} from {since.date()} to now...")
    asyncio.run(_run(channel_id, since, Path(args.out)))


if __name__ == "__main__":
    main()
